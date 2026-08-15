"""DOCX document exporter for regulatory safety reports."""

from pathlib import Path
from typing import List, Union
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from genar.models.report import ReportDocument


def _set_cell_background(cell, color_hex: str):
    """Set background shading color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell margins in dxa units."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def export_docx_report(report: ReportDocument, output_path: Union[str, Path]) -> Path:
    """Generate and export a styled Microsoft Word (.docx) regulatory safety report."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # Configure document margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. Main Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"Periodic Adverse Drug Experience Report (PADER)")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = Pt(4)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run(f"Product: {report.metadata.product_name} | FDA 21 CFR 314.80")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(71, 85, 105)
    subtitle_p.paragraph_format.space_after = Pt(16)

    # 2. Metadata Summary Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Report Identifier:", report.metadata.report_id),
        ("Manufacturer / MAH:", report.metadata.manufacturer),
        ("Reporting Period:", f"{report.metadata.reporting_period_start} to {report.metadata.reporting_period_end}"),
        ("Execution Run ID / Version:", f"{report.metadata.run_id} (v{report.metadata.app_version})"),
        ("Compliance Status:", "APPROVED - Verification Complete" if report.metadata.is_fully_approved else "PENDING REVIEW"),
    ]
    for row_idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        _set_cell_background(cell_lbl, "F1F5F9")
        _set_cell_background(cell_val, "FFFFFF")
        _set_cell_margins(cell_lbl, 80, 80, 120, 120)
        _set_cell_margins(cell_val, 80, 80, 120, 120)

        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(10)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(51, 65, 85)

        p_val = cell_val.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.name = "Calibri"
        r_val.font.size = Pt(10)
        r_val.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. Render Report Sections
    for section in report.sections:
        lines = section.markdown_content.strip().split("\n")
        in_table = False
        table_lines: List[str] = []

        for line in lines:
            line_str = line.strip()

            # Heading 2 (Section Title)
            if line_str.startswith("## "):
                h2_text = line_str.replace("## ", "").strip()
                h_p = doc.add_paragraph()
                h_run = h_p.add_run(h2_text)
                h_run.font.name = "Calibri"
                h_run.font.size = Pt(15)
                h_run.font.bold = True
                h_run.font.color.rgb = RGBColor(30, 41, 59)
                h_p.paragraph_format.space_before = Pt(14)
                h_p.paragraph_format.space_after = Pt(6)
                continue

            # Heading 3 (Subheading)
            if line_str.startswith("### "):
                h3_text = line_str.replace("### ", "").strip()
                h_p = doc.add_paragraph()
                h_run = h_p.add_run(h3_text)
                h_run.font.name = "Calibri"
                h_run.font.size = Pt(12)
                h_run.font.bold = True
                h_run.font.color.rgb = RGBColor(51, 65, 85)
                h_p.paragraph_format.space_before = Pt(10)
                h_p.paragraph_format.space_after = Pt(4)
                continue

            # Table row parsing
            if line_str.startswith("|") and line_str.endswith("|"):
                if not in_table:
                    in_table = True
                    table_lines = []
                table_lines.append(line_str)
                continue
            else:
                if in_table:
                    _render_markdown_table_in_docx(doc, table_lines)
                    in_table = False
                    table_lines = []

            if not line_str or line_str == "---":
                continue

            # Regular Narrative Paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            # Simple bold formatting parser
            parts = line_str.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.name = "Calibri"
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                if i % 2 == 1:
                    run.font.bold = True

        if in_table:
            _render_markdown_table_in_docx(doc, table_lines)

    # 4. Review & Sign-Off Block
    doc.add_page_break()
    sign_h = doc.add_paragraph()
    sign_run = sign_h.add_run("Regulatory Review & Sign-Off Record")
    sign_run.font.name = "Calibri"
    sign_run.font.size = Pt(15)
    sign_run.font.bold = True
    sign_run.font.color.rgb = RGBColor(30, 41, 59)

    sign_table = doc.add_table(rows=len(report.review_records) + 1, cols=4)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Section ID", "Status", "Reviewer Name", "Timestamp (UTC)"]
    
    # Header Row
    for col_idx, h_text in enumerate(headers):
        cell = sign_table.rows[0].cells[col_idx]
        _set_cell_background(cell, "0F172A")
        _set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, rec in enumerate(report.review_records, start=1):
        row_cells = sign_table.rows[row_idx].cells
        ts_str = rec.reviewed_at.strftime("%Y-%m-%d %H:%M:%S") if rec.reviewed_at else "Pending"
        row_vals = [rec.section_id, rec.status.value, rec.reviewer_name or "Pending", ts_str]
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"

        for col_idx, val_text in enumerate(row_vals):
            cell = row_cells[col_idx]
            _set_cell_background(cell, bg_color)
            _set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val_text)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 41, 59)

    doc.save(str(output_path))
    return output_path


def _render_markdown_table_in_docx(doc: docx.Document, table_lines: List[str]):
    """Convert a set of markdown table lines into a native Word table."""
    parsed_rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip("|").split("|")]
        # Skip separator row
        if all(set(c).issubset({"-", ":", " "}) for c in cells):
            continue
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    num_rows = len(parsed_rows)
    num_cols = max(len(r) for r in parsed_rows)

    word_table = doc.add_table(rows=num_rows, cols=num_cols)
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row_data in enumerate(parsed_rows):
        is_header = (row_idx == 0)
        row = word_table.rows[row_idx]
        bg = "1E293B" if is_header else ("F8FAFC" if row_idx % 2 == 0 else "FFFFFF")

        for col_idx in range(num_cols):
            cell = row.cells[col_idx]
            val = row_data[col_idx] if col_idx < len(row_data) else ""
            _set_cell_background(cell, bg)
            _set_cell_margins(cell, 80, 80, 120, 120)

            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5 if is_header else 9)
            if is_header:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
